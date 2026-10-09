from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

slides_data = [
    {
        "title": "Intelligent AI Food Demand Prediction Platform",
        "subtitle": "Fall 2026-27 Review II: Requirements Analysis, Architecture & 20% Prototype Milestone",
        "bullets": [
            "Team: B.Tech CSE (AI/ML), VIT Chennai",
            "Target: Optimizing Daily Kitchen Preparation & Preventing Food Waste",
            "Focus: Multi-Factor Machine Learning & Decision Analytics Engine"
        ]
    },
    {
        "title": "Problem Statement & Justification",
        "subtitle": "Addressing Food Waste in Campus & Commercial Canteens",
        "bullets": [
            "Problem: Fluctuating daily demand causes severe food waste or food shortages.",
            "Operational Pain Point: Manual estimation relies on static guesswork.",
            "Core Question Targeted: 'How much food should we prepare today?'",
            "Economic & Environmental Impact: Food waste leads to direct financial loss and elevated carbon emissions (2.5 kg CO2e per kg wasted)."
        ]
    },
    {
        "title": "System Requirements Analysis",
        "subtitle": "Functional & Non-Functional Project Scope",
        "bullets": [
            "Functional Requirement 1: Ingest multi-variable context (Historical logs, headcount, weather, calendar).",
            "Functional Requirement 2: Predict exact demand quantities (Kg) using ensemble ML algorithms.",
            "Functional Requirement 3: Provide Explainable AI (XAI) output drivers to build kitchen staff trust.",
            "Functional Requirement 4: Calculate financial risk (₹) and landfill carbon footprint (CO2e).",
            "Non-Functional Requirements: Response time < 300ms, modular API backend, intuitive visual interface."
        ]
    },
    {
        "title": "System Architecture & Operational Pipeline",
        "subtitle": "End-to-End Data & Decision Flow",
        "bullets": [
            "Layer 1 (Data Ingestion): Historical consumption logs + real-time environmental context.",
            "Layer 2 (ML Inference): Random Forest Regressor engine trained on non-linear demand patterns.",
            "Layer 3 (Decision & XAI Support): Operational buffer calculation (+5%) + rule-based feature drivers.",
            "Layer 4 (Presentation UI): Streamlit interactive dashboard with real-time Plotly gauges."
        ]
    },
    {
        "title": "Tool & Technology Stack Selection",
        "subtitle": "Justification for Technical Choices",
        "bullets": [
            "Frontend & Prototyping: Streamlit (Rapid, interactive dashboarding for live panel review).",
            "Machine Learning Framework: Scikit-Learn (Random Forest Regressor for robust tabular forecasting).",
            "Data Processing: Pandas & NumPy (Time-series aggregation and feature engineering).",
            "Data Visualization: Plotly (Dynamic gauge tracking prep thresholds vs. predicted demand)."
        ]
    },
    {
        "title": "20% Prototype Demonstration & Key Features",
        "subtitle": "Delivering an End-to-End Tracer Bullet System",
        "bullets": [
            "Multi-Variable Prediction: Ingests past consumption, expected headcount, day type, and weather forecast.",
            "Real-Time Analytics: Instant calculation of demand, surplus risk, shortage deficits, and safety buffers.",
            "Dual-Impact Metrics: Converts food waste volume into financial loss (₹) and CO2e emissions.",
            "Explainable AI (XAI): Generates natural language summaries explaining prediction drivers."
        ]
    },
    {
        "title": "Project Roadmap & Upcoming Milestones",
        "subtitle": "Plan for Review III (30% Target) & Beyond",
        "bullets": [
            "Review III Target (Oct 2026 - 30%): Automated weather API integration + benchmarking XGBoost & LightGBM.",
            "Review IV Target (Jan 2027 - 50%): PostgreSQL backend + FastAPI REST microservice integration.",
            "Review V Target (Mar 2027 - 80%): SHAP/LIME integration + field deployment in campus canteen.",
            "Review VI Target (Apr 2027 - 100%): Final Open House presentation & full report submission."
        ]
    }
]

for data in slides_data:
    slide = prs.slides.add_slide(blank_layout)
    
    # Header shape
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p_title = tf.paragraphs[0]
    p_title.text = data["title"]
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(15, 34, 64)
    
    p_sub = tf.add_paragraph()
    p_sub.text = data["subtitle"]
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    # Content box
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf_c = content_box.text_frame
    tf_c.word_wrap = True
    
    for i, bullet in enumerate(data["bullets"]):
        p = tf_c.add_paragraph() if i > 0 else tf_c.paragraphs[0]
        p.text = "• " + bullet
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(30, 30, 30)
        p.space_after = Pt(12)

prs.save("Review_2_Intelligent_Food_Demand_AI.pptx")
print("✅ PowerPoint presentation created successfully: Review_2_Intelligent_Food_Demand_AI.pptx")